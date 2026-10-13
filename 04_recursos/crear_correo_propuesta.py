# -*- coding: utf-8 -*-
"""Genera el correo de envío de la Propuesta del Proyecto 2 (docx + txt)."""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "03_entregables")

ASUNTO = "Propuesta Proyecto 2 — Visualización I · Equipo 7 (Trujillo y Figueroa)"

CUERPO = """Estimadas Adriana y Paola:

Reciban un cordial saludo. Somos el Equipo 7 (Joshua Trujillo y Nicolás Figueroa) del curso de Visualización I, y les enviamos para su aprobación la propuesta del Proyecto 2 «Visualización reproducible en Python», conforme al paso 1 del enunciado (audiencia técnica concreta + pregunta analítica).

PROPUESTA SINTÉTICA

1. Audiencia técnica: unidad de investigación comparativa sobre determinantes del bienestar de una organización internacional (p. ej. el grupo de investigación sobre desarrollo humano y calidad de vida del PNUD o del Banco Mundial), que necesita identificar de manera fundamentada y reproducible qué combinaciones de condiciones se asocian con altos niveles de bienestar reportado.

2. Pregunta analítica: ¿Qué «recetas» de bienestar predominan entre los países con mayor y menor evaluación de vida según el World Happiness Report 2024, y en qué medida el componente no explicado (Distopía + residual) altera o cuestiona esas recetas?

3. Subpreguntas: (i) qué combinación de los seis indicadores caracteriza a los terciles superior e inferior del ladder score; (ii) qué países presentan el componente Dystopia + residual más alto y qué tienen en común; (iii) si los países atípicos respecto de su receta esperada concentran residual alto o bajo.

4. Datos y variables: archivo oficial WHR24_Data_Figure_2.1.xls del World Happiness Report 2024 (143 países). Se usan el ladder score, los seis indicadores explicativos (PIB per cápita, apoyo social, esperanza de vida saludable, libertad, generosidad y percepción de corrupción) y el componente no explicado Dystopia + residual; los intervalos superior/inferior se conservan como bandas de incertidumbre. No se crearán indicadores derivados salvo la clasificación en terciles del ladder score (pandas.qcut, q=3), documentada en el notebook.

5. Plan de visualizaciones: cuatro estáticas con Seaborn (distribución de ladder score y residual; cajas de los seis indicadores por tercil; mapa de calor de composición relativa; residual frente a ladder score) y dos interactivas con Plotly (dispersión con tamaño por residual y anotación de países; barras top-15 con selector de indicador y color por tercil), cada una con su interpretación escrita.

6. Coherencia: audiencia → pregunta → variables → preparación (renombre documentado, terciles y manejo de 3 países con faltantes) → visualizaciones → interpretación tras cada gráfico → conclusiones sustentadas en la evidencia. Se declara que las asociaciones observadas no implican causalidad.

Les adjuntamos la propuesta completa en PDF (Equipo7_Proyecto2_Propuesta.pdf) y, si lo prefieren, la versión en Word. Quedamos atentos a sus observaciones y a la aprobación de la propuesta en la sesión siguiente.

Cordialmente,

Joshua Trujillo
Nicolás Figueroa
Equipo 7 — Visualización I
Quinto Semestre de 2026
Universidad Externado de Colombia
"""

ADJUNTOS = [
    "Equipo7_Proyecto2_Propuesta.pdf   (propuesta completa, versión recomendada)",
    "Equipo7_Proyecto2_Propuesta.docx  (opcional, versión editable)",
]


def build_docx(path: str) -> None:
    doc = Document()
    est = doc.styles["Normal"]
    est.font.name = "Calibri"
    est.font.size = Pt(11)

    h = doc.add_heading("Correo para el envío de la Propuesta — Proyecto 2", level=1)

    meta = doc.add_paragraph()
    meta.add_run("Curso: ").bold = True
    meta.add_run("Visualización I · Quinto Semestre 2026 · Universidad Externado de Colombia\n")
    meta.add_run("Equipo: ").bold = True
    meta.add_run("Equipo 7 — Joshua Trujillo y Nicolás Figueroa\n")
    meta.add_run("Docentes: ").bold = True
    meta.add_run("Adriana Pineda y Paola Moreno\n")
    meta.add_run("Destinatarios: ").bold = True
    meta.add_run("adriana.pineda@externado.edu.co; paola.moreno@externado.edu.co  (confirmar direcciones)")

    p = doc.add_paragraph()
    p.add_run("Asunto: ").bold = True
    r = p.add_run(ASUNTO)
    r.bold = True

    doc.add_paragraph("Adjuntos: " + ADJUNTOS[0])
    doc.paragraphs[-1].runs[0].font.size = Pt(10)
    p2 = doc.add_paragraph(ADJUNTOS[1])
    p2.runs[0].font.size = Pt(10)

    doc.add_paragraph("")  # espacio

    for bloque in CUERPO.strip().split("\n\n"):
        parrafo = doc.add_paragraph()
        if bloque.startswith("PROPUESTA SINTÉTICA") or bloque.startswith("Cordialmente"):
            parrafo.add_run(bloque).bold = True
        else:
            parrafo.add_run(bloque)
        parrafo.paragraph_format.space_after = Pt(8)

    doc.save(path)


def build_txt(path: str) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("ASUNTO: " + ASUNTO + "\n")
        fh.write("PARA: adriana.pineda@externado.edu.co; paola.moreno@externado.edu.co\n")
        fh.write("ADJUNTOS: " + "; ".join(a.split("(")[0].strip() for a in ADJUNTOS) + "\n")
        fh.write("\n" + "-" * 70 + "\n\n")
        fh.write(CUERPO)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    build_docx(os.path.join(OUT, "Correo_Propuesta_Proyecto2_Equipo7.docx"))
    build_txt(os.path.join(OUT, "Correo_Propuesta_Proyecto2_Equipo7.txt"))
    print("OK")
