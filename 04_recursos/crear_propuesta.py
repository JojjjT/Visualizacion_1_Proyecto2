# -*- coding: utf-8 -*-
"""Genera el DOCX de la propuesta del Proyecto 2 (Equipo 7)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GREEN = RGBColor(0x2E, 0x5E, 0x4E)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()

# Márgenes
for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)

# ---- Portada / encabezado
t = doc.add_heading("Propuesta — Proyecto 2", level=0)
for r in t.runs:
    r.font.color.rgb = GREEN

sub = doc.add_paragraph()
r = sub.add_run("Visualización reproducible en Python · Visualización I · Quinto Semestre 2026")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = BLUE

meta = doc.add_paragraph()
meta.add_run("Equipo: ").bold = True
meta.add_run("Equipo 7 — Joshua Trujillo y Nicolás Figueroa")
meta.add_run("    |    ").font.color.rgb = RGBColor(0x99, 0x99, 0x99)
meta.add_run("Fecha: ").bold = True
meta.add_run("12 de octubre de 2026")
meta2 = doc.add_paragraph()
meta2.add_run("Docentes: ").bold = True
meta2.add_run("Adriana Pineda – Paola Moreno")
meta2.add_run("    |    ").font.color.rgb = RGBColor(0x99, 0x99, 0x99)
meta2.add_run("Entrega del proyecto: ").bold = True
meta2.add_run("26 de octubre de 2026")

# (se elimina el separador de guiones: LaTeX lo compila con overfull)

# ---- 1. Audiencia
doc.add_heading("1. Audiencia técnica específica", level=1)

p = doc.add_paragraph()
p.add_run("Destinatario: ").bold = True
p.add_run(
    "unidad de investigación comparativa sobre determinantes del bienestar de una "
    "organización internacional (por ejemplo, el grupo de investigación del PNUD o del "
    "Banco Mundial que estudia indicadores de desarrollo y calidad de vida)."
)

p = doc.add_paragraph()
p.add_run("Interés en la problemática: ").bold = True
p.add_run(
    "este equipo necesita identificar de manera fundamentada y reproducible qué "
    "combinaciones de condiciones —PIB per cápita, apoyo social, esperanza de vida "
    "saludable, libertad, generosidad y percepción de corrupción— se asocian con altos "
    "niveles de bienestar reportado, y en qué medida esas condiciones no alcanzan a "
    "explicar la evaluación de vida de algunos países. El resultado orienta futuras "
    "investigaciones sobre los determinantes explicados y no explicados del bienestar y "
    "sirve como base verificable para estudios comparados."
)

# ---- 2. Pregunta analítica
doc.add_heading("2. Pregunta analítica", level=1)

q = doc.add_paragraph()
q.paragraph_format.left_indent = Cm(0.5)
q.paragraph_format.right_indent = Cm(0.5)
qr = q.add_run(
    "¿Qué «recetas» de bienestar predominan entre los países con mayor y menor "
    "evaluación de vida según el World Happiness Report 2024, y en qué medida el "
    "componente no explicado (Distopía + residual) altera o cuestiona esas recetas?"
)
qr.bold = True
qr.font.size = Pt(11.5)
qr.font.color.rgb = GREEN

doc.add_paragraph("Subpreguntas de apoyo:", style="Intense Quote")
subs = [
    "¿Qué combinación de los seis indicadores explicativos caracteriza a los países en "
    "los terciles superior e inferior del ladder score?",
    "¿Qué países presentan el componente Dystopia + residual más alto y qué tienen en "
    "común, sugiriendo que las variables disponibles no bastan para explicar su bienestar?",
    "¿Los países atípicos respecto de su «receta» esperada concentran altos o bajos "
    "niveles de residual, lo que indicaría que la explicación está incompleta?",
]
for s in subs:
    doc.add_paragraph(s, style="List Number")

# ---- 3. Variables
doc.add_heading("3. Variables que se emplearán", level=1)

rows = [
    ("Country name", "País", "Identificación y anotaciones en gráficos"),
    ("Ladder score", "Evaluación de vida (0–10)", "Variable de interés / eje de orden"),
    ("upperwhisker / lowerwhisker", "Intervalos del ladder score", "Bandas de incertidumbre (referencia)"),
    ("Explained by: Log GDP per capita", "Contribución del PIB per cápita", "Dimensión 1 de la «receta»"),
    ("Explained by: Social support", "Contribución del apoyo social", "Dimensión 2 de la «receta»"),
    ("Explained by: Healthy life expectancy", "Contribución de la esperanza de vida saludable", "Dimensión 3 de la «receta»"),
    ("Explained by: Freedom to make life choices", "Contribución de la libertad", "Dimensión 4 de la «receta»"),
    ("Explained by: Generosity", "Contribución de la generosidad", "Dimensión 5 de la «receta»"),
    ("Explained by: Perceptions of corruption", "Contribución de la percepción de corrupción", "Dimensión 6 de la «receta»"),
    ("Dystopia + residual", "Componente residual no explicado", "Eje central de la subpregunta 2"),
]

table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
for i, h in enumerate(["Variable (archivo oficial)", "Significado", "Uso previsto"]):
    hdr[i].text = ""
    rr = hdr[i].paragraphs[0].add_run(h)
    rr.bold = True
for var, sig, uso in rows:
    cells = table.add_row().cells
    cells[0].text = var
    cells[1].text = sig
    cells[2].text = uso

# anchos aproximados
for row in table.rows:
    row.cells[0].width = Cm(6.2)
    row.cells[1].width = Cm(5.2)
    row.cells[2].width = Cm(5.4)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Criterio de selección: ").bold = True
p.add_run(
    "se emplean las variables numéricas disponibles en el archivo oficial "
    "WHR24_Data_Figure_2.1.xls. No se crearán indicadores derivados salvo la "
    "clasificación en terciles del ladder score (pandas qcut, q=3), cuyo criterio se "
    "documentará en el notebook. Los tres países con valores faltantes (Bahrain, "
    "Tajikistan y State of Palestine) se conservarán y se excluirán explícitamente solo "
    "en las visualizaciones que lo requieran."
)

# ---- 4. Plan de visualizaciones
doc.add_heading("4. Plan de visualizaciones (anteproyecto)", level=1)

plan = [
    ("Seaborn 1", "Distribución del ladder score y del residual", "Forma global y extremos"),
    ("Seaborn 2", "Boxplot de los 6 indicadores por tercil de bienestar", "Comparación de «recetas»"),
    ("Seaborn 3", "Heatmap: composición relativa (%) de indicadores por grupo", "Perfil de receta por grupo"),
    ("Seaborn 4", "Residual (Distopía + residual) vs. ladder score", "Qué queda sin explicar"),
    ("Plotly 1", "Scatter interactivo: indicadores vs. ladder, tamaño = residual", "Explorar atípicos con hover"),
    ("Plotly 2", "Barras top-15 con selector de indicador y color por tercil", "Profundizar por dimensión"),
]
t2 = doc.add_table(rows=1, cols=3)
t2.style = "Light Grid Accent 1"
h2 = t2.rows[0].cells
for i, h in enumerate(["#", "Visualización", "Propósito analítico"]):
    h2[i].text = ""
    rr = h2[i].paragraphs[0].add_run(h)
    rr.bold = True
for a, b, c in plan:
    cells = t2.add_row().cells
    cells[0].text = a
    cells[1].text = b
    cells[2].text = c
for row in t2.rows:
    row.cells[0].width = Cm(2.6)
    row.cells[1].width = Cm(8.4)
    row.cells[2].width = Cm(5.8)

doc.add_paragraph()

# ---- 5. Coherencia
doc.add_heading("5. Coherencia esperada", level=1)
doc.add_paragraph(
    "audiencia (unidad de investigación comparativa) ⇒ "
    "pregunta (recetas de bienestar y componente no explicado) ⇒ "
    "variables (seis indicadores + residual + ladder score) ⇒ "
    "preparación (renombre documentado, terciles, manejo de faltantes) ⇒ "
    "visualizaciones (composición por grupos, perfiles, residual, atípicos) ⇒ "
    "interpretación escrita tras cada gráfico ⇒ "
    "conclusiones (qué recetas predominan y qué queda sin explicar)."
)

doc.add_paragraph()
est = doc.add_paragraph()
est.add_run("Estado: ").bold = True
est.add_run("☐ Borrador    ☑ Enviado para aprobación    ☐ Aprobado")

out = r"VISUALIZACION 1\Proyectos\Proyecto_2\03_entregables\Equipo7_Proyecto2_Propuesta.docx"
doc.save(out)
print("Guardado:", out)
